from __future__ import annotations

import re
import tempfile
import zipfile
from pathlib import Path
from typing import Callable

from openpyxl import load_workbook

_ARGB_HEX_RE = re.compile(r"^[0-9A-Fa-f]{8}$")
_RGB_HEX_RE = re.compile(r"^[0-9A-Fa-f]{6}$")
_RGB_FUNC_RE = re.compile(
    r"^(?:(?P<alpha>[0-9A-Fa-f]{2}))?RGB\(\s*(?P<r>\d{1,3})\s*,\s*(?P<g>\d{1,3})\s*,\s*(?P<b>\d{1,3})\s*\)$",
    re.IGNORECASE,
)
_RGB_ATTR_RE = re.compile(r'rgb="([^"]+)"')


def _coerce_argb(value: str) -> str:
    if _ARGB_HEX_RE.fullmatch(value):
        return value.upper()

    if _RGB_HEX_RE.fullmatch(value):
        return f"FF{value.upper()}"

    match = _RGB_FUNC_RE.fullmatch(value)
    if match:
        alpha = (match.group("alpha") or "FF").upper()
        r = max(0, min(255, int(match.group("r"))))
        g = max(0, min(255, int(match.group("g"))))
        b = max(0, min(255, int(match.group("b"))))
        return f"{alpha}{r:02X}{g:02X}{b:02X}"

    # Last-resort fallback for malformed values that openpyxl would reject.
    return "FF000000"


def _sanitize_stylesheet(styles_xml: bytes) -> tuple[bytes, int]:
    text = styles_xml.decode("utf-8", errors="replace")
    replacements = 0

    def repl(match: re.Match[str]) -> str:
        nonlocal replacements
        original = match.group(1)
        fixed = _coerce_argb(original)
        if fixed != original:
            replacements += 1
        return f'rgb="{fixed}"'

    updated = _RGB_ATTR_RE.sub(repl, text)
    return updated.encode("utf-8"), replacements


def _create_sanitized_copy(input_path: Path) -> tuple[Path, int]:
    with zipfile.ZipFile(input_path) as src:
        styles_name = "xl/styles.xml"
        if styles_name not in src.namelist():
            raise ValueError("Workbook is missing xl/styles.xml")

        sanitized_styles, replacements = _sanitize_stylesheet(src.read(styles_name))
        temp = tempfile.NamedTemporaryFile(prefix="xlsx-style-fix-", suffix=".xlsx", delete=False)
        temp_path = Path(temp.name)
        temp.close()

        with zipfile.ZipFile(temp_path, "w") as dst:
            for info in src.infolist():
                data = sanitized_styles if info.filename == styles_name else src.read(info.filename)
                dst.writestr(info, data)

    return temp_path, replacements


def load_workbook_compat(input_path: Path, log: Callable[[str], None] | None = None):
    try:
        return load_workbook(input_path)
    except ValueError as exc:
        message = str(exc)
        if "could not read stylesheet" not in message:
            raise

        temp_path, replacements = _create_sanitized_copy(input_path)
        if log is not None:
            log(
                "Detected invalid workbook color styles; "
                f"sanitized {replacements} rgb value(s) in styles.xml and retried."
            )

        try:
            return load_workbook(temp_path)
        finally:
            temp_path.unlink(missing_ok=True)
