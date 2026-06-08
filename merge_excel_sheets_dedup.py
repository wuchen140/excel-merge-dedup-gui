#!/usr/bin/env python3
"""
Merge all sheets in an Excel file into one sheet and remove duplicates.

Default behavior:
- Reads first 3 columns as: 英文 / 日文 / 韩文
- If a sheet uses 英文 / 韩文 / 日文, it auto-swaps to 英文 / 日文 / 韩文
- Skips empty rows
- Deduplicates by all 3 columns (trimmed text)

Usage:
  python3 merge_excel_sheets_dedup.py /path/to/input.xlsx
  python3 merge_excel_sheets_dedup.py /path/to/input.xlsx -o /path/to/output.xlsx
  python3 merge_excel_sheets_dedup.py /path/to/input.xlsx --dedup-by en
"""

from __future__ import annotations

import argparse
from pathlib import Path
from typing import Dict, List, Tuple

from openpyxl import Workbook

from workbook_compat import load_workbook_compat

TARGET_HEADERS = ("英文", "日文", "韩文")


def norm(v) -> str:
    if v is None:
        return ""
    if isinstance(v, str):
        return v.strip()
    return str(v).strip()


def detect_col_map(ws) -> Tuple[int, int, int]:
    """
    Return (en_col, ja_col, ko_col), 1-based indices.
    Priority:
      1) Exact header match by text in row 1.
      2) Fallback to first 3 columns as (1,2,3).
    """
    header_map: Dict[str, int] = {}
    for c in range(1, max(ws.max_column, 3) + 1):
        h = norm(ws.cell(1, c).value)
        if h in TARGET_HEADERS and h not in header_map:
            header_map[h] = c

    if all(h in header_map for h in TARGET_HEADERS):
        return header_map["英文"], header_map["日文"], header_map["韩文"]

    # fallback
    return 1, 2, 3


def output_path_from_input(input_path: Path) -> Path:
    return input_path.with_name(f"{input_path.stem}_合并去重{input_path.suffix}")


def main() -> int:
    parser = argparse.ArgumentParser(description="Merge Excel sheets and deduplicate rows.")
    parser.add_argument("input", help="Input .xlsx path")
    parser.add_argument("-o", "--output", help="Output .xlsx path")
    parser.add_argument(
        "--dedup-by",
        choices=["all", "en"],
        default="all",
        help="Deduplicate by all 3 columns (default) or only English column",
    )
    parser.add_argument("--sheet-name", default="合并去重", help="Output sheet name")

    args = parser.parse_args()

    input_path = Path(args.input).expanduser().resolve()
    if not input_path.exists():
        raise FileNotFoundError(f"Input file not found: {input_path}")
    if input_path.suffix.lower() != ".xlsx":
        raise ValueError("Only .xlsx is supported")

    out_path = Path(args.output).expanduser().resolve() if args.output else output_path_from_input(input_path)

    wb = load_workbook_compat(input_path, log=print)

    merged_rows: List[Tuple[str, str, str]] = []
    seen = set()

    total_read = 0
    total_empty = 0
    total_dup = 0

    for ws in wb.worksheets:
        en_col, ja_col, ko_col = detect_col_map(ws)

        sheet_read = 0
        sheet_empty = 0
        sheet_dup = 0
        sheet_keep = 0

        for r in range(2, ws.max_row + 1):
            en = norm(ws.cell(r, en_col).value)
            ja = norm(ws.cell(r, ja_col).value)
            ko = norm(ws.cell(r, ko_col).value)

            if not en and not ja and not ko:
                total_empty += 1
                sheet_empty += 1
                continue

            total_read += 1
            sheet_read += 1

            key = (en, ja, ko) if args.dedup_by == "all" else (en,)
            if key in seen:
                total_dup += 1
                sheet_dup += 1
                continue

            seen.add(key)
            merged_rows.append((en, ja, ko))
            sheet_keep += 1

        print(
            f"[{ws.title}] read={sheet_read}, keep={sheet_keep}, "
            f"empty={sheet_empty}, duplicate={sheet_dup}, "
            f"col_map=(EN:{en_col}, JA:{ja_col}, KO:{ko_col})"
        )

    wb_out = Workbook()
    ws_out = wb_out.active
    ws_out.title = args.sheet_name
    ws_out.append(list(TARGET_HEADERS))

    for row in merged_rows:
        ws_out.append(list(row))

    wb_out.save(out_path)

    print("\nDone")
    print(f"Input : {input_path}")
    print(f"Output: {out_path}")
    print(f"Rows read (non-empty): {total_read}")
    print(f"Rows kept          : {len(merged_rows)}")
    print(f"Rows removed dup   : {total_dup}")
    print(f"Rows skipped empty : {total_empty}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
