#!/usr/bin/env python3
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Dict, List, Tuple

from PySide6.QtCore import QObject, QThread, Signal
from PySide6.QtWidgets import (
    QApplication,
    QButtonGroup,
    QFileDialog,
    QFrame,
    QGridLayout,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QPlainTextEdit,
    QVBoxLayout,
    QWidget,
)

TARGET_HEADERS = ("英文", "日文", "韩文")


def norm(v) -> str:
    if v is None:
        return ""
    if isinstance(v, str):
        return v.strip()
    return str(v).strip()


def detect_col_map(ws) -> Tuple[int, int, int]:
    header_map: Dict[str, int] = {}
    for c in range(1, max(ws.max_column, 3) + 1):
        h = norm(ws.cell(1, c).value)
        if h in TARGET_HEADERS and h not in header_map:
            header_map[h] = c

    if all(h in header_map for h in TARGET_HEADERS):
        return header_map["英文"], header_map["日文"], header_map["韩文"]
    return 1, 2, 3


def output_path_from_input(input_path: Path) -> Path:
    return input_path.with_name(f"{input_path.stem}_合并去重{input_path.suffix}")


@dataclass
class MergeStats:
    rows_read: int
    rows_kept: int
    rows_dup: int
    rows_empty: int
    output_path: Path


def merge_excel_sheets(
    input_path: Path,
    output_path: Path | None,
    dedup_by: str,
    sheet_name: str,
    log: Callable[[str], None],
) -> MergeStats:
    # Lazy import to keep app cold-start fast.
    from openpyxl import Workbook
    from workbook_compat import load_workbook_compat

    if not input_path.exists():
        raise FileNotFoundError(f"输入文件不存在: {input_path}")
    if input_path.suffix.lower() != ".xlsx":
        raise ValueError("仅支持 .xlsx 文件")

    out_path = output_path or output_path_from_input(input_path)
    wb = load_workbook_compat(input_path, log)
    merged_rows: List[Tuple[str, str, str]] = []
    seen = set()

    total_read = 0
    total_empty = 0
    total_dup = 0

    for ws in wb.worksheets:
        en_col, ja_col, ko_col = detect_col_map(ws)
        sheet_read = 0
        sheet_empty = 0
        sheet_dup = 0
        sheet_keep = 0

        for r in range(2, ws.max_row + 1):
            en = norm(ws.cell(r, en_col).value)
            ja = norm(ws.cell(r, ja_col).value)
            ko = norm(ws.cell(r, ko_col).value)

            if not en and not ja and not ko:
                total_empty += 1
                sheet_empty += 1
                continue

            total_read += 1
            sheet_read += 1

            key = (en, ja, ko) if dedup_by == "all" else (en,)
            if key in seen:
                total_dup += 1
                sheet_dup += 1
                continue

            seen.add(key)
            merged_rows.append((en, ja, ko))
            sheet_keep += 1

        log(
            f"[{ws.title}] read={sheet_read}, keep={sheet_keep}, "
            f"empty={sheet_empty}, duplicate={sheet_dup}, "
            f"col_map=(EN:{en_col}, JA:{ja_col}, KO:{ko_col})"
        )

    wb_out = Workbook()
    ws_out = wb_out.active
    ws_out.title = sheet_name or "合并去重"
    ws_out.append(list(TARGET_HEADERS))
    for row in merged_rows:
        ws_out.append(list(row))
    wb_out.save(out_path)

    return MergeStats(
        rows_read=total_read,
        rows_kept=len(merged_rows),
        rows_dup=total_dup,
        rows_empty=total_empty,
        output_path=out_path,
    )


class MergeWorker(QObject):
    finished = Signal(object)
    failed = Signal(str)
    log = Signal(str)

    def __init__(self, input_path: Path, output_path: Path | None, dedup_by: str, sheet_name: str) -> None:
        super().__init__()
        self.input_path = input_path
        self.output_path = output_path
        self.dedup_by = dedup_by
        self.sheet_name = sheet_name

    def run(self) -> None:
        try:
            stats = merge_excel_sheets(
                input_path=self.input_path,
                output_path=self.output_path,
                dedup_by=self.dedup_by,
                sheet_name=self.sheet_name,
                log=self.log.emit,
            )
            self.finished.emit(stats)
        except Exception as exc:
            self.failed.emit(str(exc))


class MainWindow(QMainWindow):
    def __init__(self) -> None:
        super().__init__()
        self.setWindowTitle("Excel 多表合并去重工具")
        self.resize(920, 640)

        self.thread: QThread | None = None
        self.worker: MergeWorker | None = None

        page = QWidget()
        root = QVBoxLayout(page)
        root.setContentsMargins(20, 20, 20, 20)
        root.setSpacing(0)

        shell = QFrame()
        shell.setObjectName("shellCard")
        shell_layout = QVBoxLayout(shell)
        shell_layout.setContentsMargins(20, 14, 20, 16)
        shell_layout.setSpacing(12)

        title = QLabel("Excel 多表合并去重工具")
        title.setObjectName("titleLabel")
        title.setStyleSheet("font-size: 24px; font-weight: 700;")
        subtitle = QLabel("选择输入文件、输出路径和去重方式，然后执行合并。")
        subtitle.setObjectName("subLabel")

        form_card = QFrame()
        form_card.setObjectName("formCard")
        form_layout = QGridLayout(form_card)
        form_layout.setContentsMargins(0, 0, 0, 0)
        form_layout.setHorizontalSpacing(8)
        form_layout.setVerticalSpacing(12)

        self.input_edit = QLineEdit()
        self.output_edit = QLineEdit()
        self.sheet_name_edit = QLineEdit("合并去重")
        self.input_btn = QPushButton("浏览")
        self.output_btn = QPushButton("保存为")
        self.run_btn = QPushButton("开始合并")
        self.clear_log_btn = QPushButton("清空日志")
        self.input_btn.setObjectName("primaryButton")
        self.run_btn.setObjectName("primaryButton")
        self.output_btn.setObjectName("secondaryButton")
        self.clear_log_btn.setObjectName("secondaryButton")

        self.input_edit.setPlaceholderText("未选择文件")
        self.output_edit.setPlaceholderText("默认：输入文件名_合并去重.xlsx")

        self.all_btn = QPushButton("全部列")
        self.en_btn = QPushButton("仅英文列")
        self.all_btn.setCheckable(True)
        self.en_btn.setCheckable(True)
        self.all_btn.setChecked(True)
        self.all_btn.setObjectName("segmentButton")
        self.en_btn.setObjectName("segmentButton")
        self.dedup_group = QButtonGroup(self)
        self.dedup_group.setExclusive(True)
        self.dedup_group.addButton(self.all_btn)
        self.dedup_group.addButton(self.en_btn)

        form_layout.addWidget(QLabel("输入文件"), 0, 0)
        form_layout.addWidget(self.input_edit, 0, 1)
        form_layout.addWidget(self.input_btn, 0, 2)

        form_layout.addWidget(QLabel("输出文件"), 1, 0)
        form_layout.addWidget(self.output_edit, 1, 1)
        form_layout.addWidget(self.output_btn, 1, 2)

        dedup_row = QWidget()
        dedup_layout = QHBoxLayout(dedup_row)
        dedup_layout.setContentsMargins(0, 0, 0, 0)
        dedup_layout.setSpacing(0)
        dedup_wrap = QFrame()
        dedup_wrap.setObjectName("segmentWrap")
        dedup_wrap_layout = QHBoxLayout(dedup_wrap)
        dedup_wrap_layout.setContentsMargins(4, 4, 4, 4)
        dedup_wrap_layout.setSpacing(4)
        dedup_wrap_layout.addWidget(self.all_btn)
        dedup_wrap_layout.addWidget(self.en_btn)
        dedup_layout.addWidget(dedup_wrap)
        dedup_layout.addStretch()

        form_layout.addWidget(QLabel("去重方式"), 2, 0)
        form_layout.addWidget(dedup_row, 2, 1, 1, 2)

        form_layout.addWidget(QLabel("工作表名"), 3, 0)
        form_layout.addWidget(self.sheet_name_edit, 3, 1, 1, 2)
        form_layout.setColumnStretch(1, 1)

        action_row = QWidget()
        action_layout = QHBoxLayout(action_row)
        action_layout.setContentsMargins(0, 0, 0, 0)
        action_layout.setSpacing(8)
        action_layout.addWidget(self.run_btn)
        action_layout.addWidget(self.clear_log_btn)
        action_layout.addStretch()

        log_card = QFrame()
        log_card.setObjectName("logCard")
        log_layout = QVBoxLayout(log_card)
        log_layout.setContentsMargins(12, 10, 12, 12)
        log_layout.setSpacing(8)
        log_title = QLabel("运行日志")
        log_title.setObjectName("logTitle")

        self.log_box = QPlainTextEdit()
        self.log_box.setObjectName("logBox")
        self.log_box.setReadOnly(True)
        self.log_box.setPlaceholderText("等待开始...")
        log_layout.addWidget(log_title)
        log_layout.addWidget(self.log_box, 1)

        shell_layout.addWidget(title)
        shell_layout.addWidget(subtitle)
        shell_layout.addWidget(form_card)
        shell_layout.addWidget(action_row)
        shell_layout.addWidget(log_card, 1)
        root.addWidget(shell)
        self.setCentralWidget(page)
        self.apply_styles()

        self.input_btn.clicked.connect(self.choose_input_file)
        self.output_btn.clicked.connect(self.choose_output_file)
        self.run_btn.clicked.connect(self.start_merge)
        self.clear_log_btn.clicked.connect(self.log_box.clear)

    def apply_styles(self) -> None:
        self.setStyleSheet(
            """
            QWidget {
                background: #ECECF1;
                color: #1D1D1F;
                font-family: -apple-system, BlinkMacSystemFont, "SF Pro Text", "PingFang SC", "Helvetica Neue";
                font-size: 13px;
            }
            QFrame#shellCard {
                background: #FFFFFF;
                border: 1px solid #D1D1D6;
                border-radius: 16px;
            }
            QLabel#titleLabel {
                font-size: 26px;
                font-weight: 700;
            }
            QLabel#subLabel {
                color: #6E6E73;
                font-size: 13px;
            }
            QLineEdit {
                background: #FFFFFF;
                border: 1px solid #D2D2D7;
                border-radius: 10px;
                padding: 8px 12px;
                min-height: 22px;
                selection-background-color: #007AFF;
            }
            QPushButton {
                border-radius: 10px;
                min-height: 28px;
                padding: 4px 14px;
            }
            QPushButton#primaryButton {
                background: #007AFF;
                color: #FFFFFF;
                border: none;
                font-weight: 600;
            }
            QPushButton#primaryButton:hover {
                background: #0A84FF;
            }
            QPushButton#primaryButton:disabled {
                background: #A8D2FF;
                color: #F3F8FF;
            }
            QPushButton#secondaryButton {
                background: #F2F2F7;
                color: #1D1D1F;
                border: 1px solid #D1D1D6;
            }
            QFrame#segmentWrap {
                background: #F2F2F7;
                border-radius: 10px;
            }
            QPushButton#segmentButton {
                background: transparent;
                border: 1px solid transparent;
                border-radius: 8px;
                color: #6E6E73;
                font-weight: 500;
                min-width: 72px;
                min-height: 24px;
                padding: 2px 10px;
            }
            QPushButton#segmentButton:checked {
                background: #FFFFFF;
                color: #1D1D1F;
                border: 1px solid #D8D8DD;
                font-weight: 600;
            }
            QFrame#logCard {
                background: #FBFBFD;
                border: 1px solid #E5E5EA;
                border-radius: 12px;
            }
            QLabel#logTitle {
                color: #6E6E73;
                font-weight: 600;
                font-size: 12px;
            }
            QPlainTextEdit#logBox {
                background: #FFFFFF;
                border: 1px solid #E5E5EA;
                border-radius: 10px;
                padding: 8px;
                color: #2C2C2E;
                selection-background-color: #B3D7FF;
            }
            """
        )

    def append_log(self, text: str) -> None:
        self.log_box.appendPlainText(text)

    def choose_input_file(self) -> None:
        file_path, _ = QFileDialog.getOpenFileName(self, "选择输入 Excel 文件", "", "Excel (*.xlsx)")
        if not file_path:
            return
        self.input_edit.setText(file_path)
        if not self.output_edit.text().strip():
            default_out = output_path_from_input(Path(file_path))
            self.output_edit.setText(str(default_out))

    def choose_output_file(self) -> None:
        current = self.output_edit.text().strip()
        default_name = Path(current).name if current else "合并去重结果.xlsx"
        default_dir = str(Path(current).parent) if current else ""
        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "选择输出文件",
            str(Path(default_dir) / default_name) if default_dir else default_name,
            "Excel (*.xlsx)",
        )
        if file_path:
            self.output_edit.setText(file_path)

    def start_merge(self) -> None:
        input_raw = self.input_edit.text().strip()
        if not input_raw:
            QMessageBox.warning(self, "缺少输入文件", "请先选择输入文件。")
            return

        input_path = Path(input_raw).expanduser().resolve()
        output_raw = self.output_edit.text().strip()
        output_path = Path(output_raw).expanduser().resolve() if output_raw else None
        sheet_name = self.sheet_name_edit.text().strip() or "合并去重"
        dedup_by = "all" if self.all_btn.isChecked() else "en"

        self.append_log("=" * 60)
        self.append_log(f"输入文件: {input_path}")
        self.append_log(f"输出文件: {output_path or '(自动生成)'}")
        self.append_log(f"去重方式: {'全部列' if dedup_by == 'all' else '仅英文列'}")
        self.append_log(f"输出工作表名: {sheet_name}")

        self.run_btn.setEnabled(False)
        self.worker = MergeWorker(input_path, output_path, dedup_by, sheet_name)
        self.thread = QThread()
        self.worker.moveToThread(self.thread)
        self.thread.started.connect(self.worker.run)
        self.worker.log.connect(self.append_log)
        self.worker.finished.connect(self.on_success)
        self.worker.failed.connect(self.on_failed)
        self.worker.finished.connect(self.cleanup_thread)
        self.worker.failed.connect(self.cleanup_thread)
        self.thread.start()

    def cleanup_thread(self) -> None:
        if self.thread is not None:
            self.thread.quit()
            self.thread.wait()
            self.thread = None
        self.worker = None
        self.run_btn.setEnabled(True)

    def on_success(self, stats: MergeStats) -> None:
        self.append_log("")
        self.append_log("Done")
        self.append_log(f"Output: {stats.output_path}")
        self.append_log(f"Rows read (non-empty): {stats.rows_read}")
        self.append_log(f"Rows kept          : {stats.rows_kept}")
        self.append_log(f"Rows removed dup   : {stats.rows_dup}")
        self.append_log(f"Rows skipped empty : {stats.rows_empty}")
        QMessageBox.information(self, "完成", f"处理完成，输出文件：\n{stats.output_path}")

    def on_failed(self, error: str) -> None:
        self.append_log(f"ERROR: {error}")
        QMessageBox.critical(self, "处理失败", error)


def main() -> int:
    app = QApplication([])
    window = MainWindow()
    window.show()
    return app.exec()


if __name__ == "__main__":
    raise SystemExit(main())
